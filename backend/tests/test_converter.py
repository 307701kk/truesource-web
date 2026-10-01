from __future__ import annotations

from datetime import date, datetime

from app.converter import KIND_DATA, KIND_SUBTOTAL, KIND_TOTAL, convert_workbook, infer_dtype

from .conftest import write_workbook


def one_sheet(tmp_path, rows, name="S", **kw):
    path = write_workbook(tmp_path / "a.xlsx", {name: rows}, **kw)
    return convert_workbook(path)[0]


def test_header_after_title_rows_and_preamble(tmp_path):
    sh = one_sheet(
        tmp_path,
        [
            ["2026년 집계"],
            ["기준일: 2026-09-28"],
            [],
            ["팀명", "월", "매출"],
            ["영업1팀", "7월", 5],
        ],
    )
    assert sh.header_row == 4
    assert sh.preamble == ["2026년 집계", "기준일: 2026-09-28"]
    assert [c.name for c in sh.columns] == ["팀명", "월", "매출"]
    assert sh.rows[0][0] == 5  # original Excel row number is kept
    assert sh.rows[0][3] == ["영업1팀", "7월", 5]


def test_header_on_first_row(tmp_path):
    sh = one_sheet(tmp_path, [["코드", "이름"], ["C001", "광명전기"], ["C002", "보광메탈"]])
    assert sh.header_row == 1
    assert sh.preamble == []
    assert sh.n_data_rows == 2


def test_row_numbers_survive_blank_rows(tmp_path):
    sh = one_sheet(tmp_path, [["a", "b"], [1, 2], [], [], [3, 4]])
    assert [r[0] for r in sh.rows] == [2, 5]


def test_subtotal_and_total_rows_are_tagged_not_dropped(tmp_path):
    sh = one_sheet(
        tmp_path,
        [
            ["일자", "업체명", "금액"],
            ["07/10", "한울산기", 10],
            [None, "한울산기 소계", 10],
            ["합계", None, 10],
        ],
    )
    kinds = [(r[1], r[2]) for r in sh.rows]
    assert kinds == [(KIND_DATA, None), (KIND_SUBTOTAL, "한울산기 소계"), (KIND_TOTAL, "합계")]
    assert sh.n_data_rows == 1
    assert sh.rows[2][3][2] == 10  # the total's number is kept


def test_column_types(tmp_path):
    sh = one_sheet(
        tmp_path,
        [
            ["i", "f", "d", "dt", "s", "m"],
            [1, 1.5, datetime(2026, 1, 2), datetime(2026, 1, 2, 3, 4), "x", 1],
            [2, 2, datetime(2026, 1, 3), datetime(2026, 1, 3, 5, 6), "y", "text"],
        ],
    )
    types = {c.name: c.dtype for c in sh.columns}
    assert types == {
        "i": "BIGINT",
        "f": "DOUBLE",
        "d": "DATE",
        "dt": "TIMESTAMP",
        "s": "VARCHAR",
        "m": "VARCHAR",
    }
    assert sh.mixed_columns == ["m"]
    assert sh.rows[0][3][2] == date(2026, 1, 2)
    assert sh.rows[0][3][5] == "1"  # mixed column stored as text


def test_total_row_label_does_not_break_numeric_column(tmp_path):
    sh = one_sheet(tmp_path, [["No", "금액"], [1, 5], [2, 6], ["합계", 11]])
    assert {c.name: c.dtype for c in sh.columns}["No"] == "BIGINT"
    assert sh.rows[2][1] == KIND_TOTAL and sh.rows[2][3][0] is None


def test_duplicate_and_missing_headers(tmp_path):
    sh = one_sheet(tmp_path, [["이름", "이름", None, "_row"], ["a", "b", "c", "d"]])
    assert [c.name for c in sh.columns] == ["이름", "이름_2", "열3", "col__row"]


def test_error_cells_become_null_and_are_counted(tmp_path):
    sh = one_sheet(tmp_path, [["a", "b"], [1, "#N/A"], [2, 3]])
    assert sh.error_cells == 1
    assert {c.name: c.dtype for c in sh.columns}["b"] == "BIGINT"
    assert sh.rows[0][3] == [1, None]


def test_hidden_empty_and_headerless_sheets(tmp_path):
    path = write_workbook(
        tmp_path / "b.xlsx",
        {
            "보임": [["a", "b"], [1, 2]],
            "숨김": [["c", "d"], [3, 4]],
            "빈시트": [],
            "숫자만": [[1, 2], [3, 4]],
        },
        hidden=("숨김",),
    )
    by_name = {s.sheet_name: s for s in convert_workbook(path)}
    assert by_name["숨김"].hidden and not by_name["보임"].hidden
    assert by_name["빈시트"].columns == [] and by_name["빈시트"].rows == []
    nohdr = by_name["숫자만"]
    assert nohdr.header_row is None and nohdr.n_data_rows == 2
    assert [c.name for c in nohdr.columns] == ["열1", "열2"]


def test_whitespace_only_cells_are_empty_and_text_is_stripped(tmp_path):
    sh = one_sheet(tmp_path, [["a", "b"], ["  x  ", "   "], ["y", "z"]])
    assert sh.rows[0][3] == ["x", None]


def test_infer_dtype_edge_cases():
    assert infer_dtype([]) == ("VARCHAR", False)
    assert infer_dtype([True, False]) == ("BOOLEAN", False)
    assert infer_dtype([1, 2.0]) == ("BIGINT", False)
    assert infer_dtype([float("inf")])[0] == "VARCHAR"
    assert infer_dtype([1, "a"]) == ("VARCHAR", True)


# ---- regressions found in code review ------------------------------------------------------
def test_duplicate_headers_differing_only_by_case_still_load_into_duckdb(tmp_path):
    from datetime import datetime as dt

    from app.store import Store

    sh = one_sheet(tmp_path, [["Name", "name", "_Row", "a\tb"], [1, 2, 3, 4]])
    assert [c.name for c in sh.columns] == ["Name", "name_2", "col__Row", "a b"]
    store = Store()
    fid = store.add_file(
        rel_path="a.xlsx",
        name="a.xlsx",
        dept="미분류",
        top_folder="",
        size=1,
        modified=dt(2026, 1, 1),
        sha256="h",
    )
    store.add_sheet(fid, sh)  # must not raise


def test_numeric_looking_header_cells(tmp_path):
    sh = one_sheet(tmp_path, [["품목", 2024, 2025, 2026], ["A", 1000, 2000, 3000]])
    assert sh.header_row == 1
    assert [c.name for c in sh.columns] == ["품목", "2024", "2025", "2026"]
    sh = one_sheet(tmp_path, [["품목", 1, 2, 3], ["A", "x", "y", "z"], ["B", 10, 20, 30]])
    assert sh.header_row == 1 and sh.n_data_rows == 2
    # a data-looking first row must still not become the header when there is a real one
    sh = one_sheet(tmp_path, [["품목", "수량", "단가"], ["A", 1, 2]])
    assert sh.header_row == 1


def test_table_not_starting_in_column_a(tmp_path):
    path = write_workbook(
        tmp_path / "off.xlsx",
        {"S": [["제목"], [], [None, None, "a", "b"], [None, None, 1, 2], [None, None, 3, 4]]},
    )
    sh = convert_workbook(path)[0]
    assert sh.header_row == 3
    assert [(c.name, c.dtype, c.source_col) for c in sh.columns] == [
        ("a", "BIGINT", 3),
        ("b", "BIGINT", 4),
    ]


def test_more_subtotal_spellings(tmp_path):
    sh = one_sheet(
        tmp_path,
        [
            ["구분", "지역", "금액"],
            ["가", "부산", 1],
            ["Subtotal", None, 1],
            ["Total", None, 1],
            ["나", "합계(원)", 1],
            [None, "계", 1],
            [None, None, "합계"],
            ["계산서", "서울", 1],  # "계산서" is data: only a bare "계" counts
        ],
    )
    kinds = [r[1] for r in sh.rows]
    assert kinds == ["data", "subtotal", "total", "total", "total", "total", "data"]


def test_decimal_subtotal_widens_integer_column(tmp_path):
    sh = one_sheet(tmp_path, [["x", "y"], [1, 1], [2, 1], ["합계", 3.5]])
    assert {c.name: c.dtype for c in sh.columns}["y"] == "DOUBLE"
    assert sh.rows[2][3][1] == 3.5


def test_styled_empty_rows_at_the_bottom_are_not_a_truncation(tmp_path):
    from openpyxl.styles import PatternFill

    path = tmp_path / "bloat.xlsx"
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    ws["A200000"].fill = PatternFill("solid", fgColor="FFFF00")
    wb.save(path)
    sh = convert_workbook(path)[0]
    assert sh.n_data_rows == 1 and sh.truncated is False


def test_row_cap_counts_only_real_rows(tmp_path, monkeypatch):
    from app import config

    monkeypatch.setattr(config, "MAX_ROWS_PER_SHEET", 3)
    sh = one_sheet(tmp_path, [["a"], [1], [2], [3], [4]] and [["a", "b"], [1, 1], [2, 2], [3, 3]])
    assert sh.truncated and len(sh.rows) == 2


def test_uncached_formulas_are_counted(tmp_path):
    from app.converter import uncached_formula_cells

    path = write_workbook(tmp_path / "f.xlsx", {"S": [["a", "b", "c"], [1, 2, "=A2+B2"]]})
    assert uncached_formula_cells(path) == 1
    plain = write_workbook(tmp_path / "p.xlsx", {"S": [["a", "b"], [1, 2]]})
    assert uncached_formula_cells(plain) == 0


def test_zip_bomb_guard(tmp_path, monkeypatch):
    import pytest

    from app import config
    from app.converter import check_archive

    path = write_workbook(tmp_path / "z.xlsx", {"S": [["a"], [1]]})
    check_archive(path)
    monkeypatch.setattr(config, "MAX_UNCOMPRESSED_BYTES", 10)
    with pytest.raises(ValueError, match="압축 해제"):
        check_archive(path)


def test_key_value_title_row_is_not_the_header(tmp_path):
    sh = one_sheet(
        tmp_path,
        [
            ["작성자: 홍길동", "부서: 영업", "단위: 원"],
            ["일자", "거래처", "금액"],
            ["2026-01-01", "가", 100],
        ],
    )
    assert sh.header_row == 2
    assert [c.name for c in sh.columns] == ["일자", "거래처", "금액"]


def test_data_first_row_with_two_numbers_is_not_a_month_header(tmp_path):
    sh = one_sheet(tmp_path, [["품목", "수량", "단가"], ["A", 1, 2], ["B", 3, 4]])
    assert sh.header_row == 1 and sh.n_data_rows == 2
    # only two numeric cells: not a month/day header, so this sheet has no header row
    sh = one_sheet(tmp_path, [["품목", "비고", 1, 2], ["A", "x", 5, 6]])
    assert sh.header_row is None


def test_total_spellings_with_punctuation(tmp_path):
    sh = one_sheet(
        tmp_path,
        [
            ["구분", "지역", "금액"],
            ["가", "부산", 1],
            ["합계:", None, 1],
            ["[합계]", None, 1],
            ["월계", None, 1],
            ["대계", None, 1],
        ],
    )
    assert [r[1] for r in sh.rows] == ["data", "total", "total", "total", "total"]
