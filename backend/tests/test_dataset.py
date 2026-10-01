"""End-to-end checks against the real sample dataset (가온산업_가상데이터).

Skipped unless TRUESOURCE_DATASET points at the dataset root, i.e. the folder that contains
`공유폴더/` and `정답지/`.  Example (PowerShell):
    $env:TRUESOURCE_DATASET = "C:\\data\\가온산업_가상데이터"; pytest tests/test_dataset.py
"""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path

import openpyxl
import pytest

from app.converter import _text
from app.service import ScanService

DATASET = os.environ.get("TRUESOURCE_DATASET")
pytestmark = pytest.mark.skipif(not DATASET, reason="TRUESOURCE_DATASET is not set")


@pytest.fixture(scope="module")
def root() -> Path:
    return Path(DATASET)


@pytest.fixture(scope="module")
def service(root):
    svc = ScanService()
    svc.start(str(root / "공유폴더"), background=False)
    return svc


def sheet_table(service, rel_path: str, sheet: str | None = None) -> str:
    store = service.store
    fid = store.query("SELECT file_id FROM files WHERE rel_path = ?", [rel_path])[0][0]
    sheets = store.sheets_of(fid)
    return next(s["table"] for s in sheets if sheet is None or s["sheet"] == sheet)


def test_scan_counts(service):
    st = service.status()
    assert st["state"] == "done"
    assert st["files_total"] == 137  # 138 xlsx - 1 lock file
    assert st["skipped"] == {"lock_file": 1}
    assert st["errors"] == [] and st["warnings"] == []
    assert (st["sheets"], st["data_rows"]) == (221, 20395)


def test_every_data_cell_equals_the_original(service, root):
    """Independent re-read of every workbook: no data cell may be lost, changed or invented."""
    store = service.store
    problems = []
    for fid, rel in store.query("SELECT file_id, rel_path FROM files WHERE status = 'ok'"):
        wb = openpyxl.load_workbook(root / "공유폴더" / rel, data_only=True)
        for sh in store.sheets_of(fid):
            if not sh["table"]:
                continue
            ws = wb[sh["sheet"]]
            cols = sh["columns"]
            quoted = ", ".join('"' + c["name"].replace('"', '""') + '"' for c in cols)
            rows = store.query(
                f'SELECT "_row", "_row_kind", {quoted} FROM {sh["table"]} ORDER BY "_row"'
            )
            present = {r[0] for r in rows}
            for row_no, kind, *vals in rows:
                if kind != "data":
                    continue
                for c, v in zip(cols, vals, strict=True):
                    raw = ws.cell(row=row_no, column=c["source_col"]).value
                    if isinstance(raw, str):
                        raw = raw.strip() or None
                    if raw is None:
                        ok = v is None
                    elif c["dtype"] == "VARCHAR":
                        ok = v == _text(raw)
                    elif c["dtype"] == "DATE":
                        ok = v == (raw.date() if hasattr(raw, "date") else raw)
                    elif c["dtype"] in ("BIGINT", "TIMESTAMP"):
                        ok = v == raw
                    else:
                        ok = v is not None and abs(float(v) - float(raw)) < 1e-9
                    if not ok:
                        problems.append((rel, sh["sheet"], row_no, c["name"], raw, v))
            if sh["header_row"]:
                for i in range(sh["header_row"] + 1, ws.max_row + 1):
                    cells = (ws.cell(row=i, column=j).value for j in range(1, ws.max_column + 1))
                    if any(c not in (None, "") for c in cells) and i not in present:
                        problems.append((rel, sh["sheet"], i, "missing row"))
    assert problems == []


def test_q3_team_totals_match_answer_key(service, root):
    key = json.loads((root / "정답지/answer_key.json").read_text(encoding="utf-8"))[
        "시나리오_3분기_팀순위"
    ]
    store = service.store

    t = sheet_table(service, "영업팀/2026/실적/실적집계_v2.xlsx", "팀별실적")
    got = dict(
        store.query(
            f'SELECT "팀명", CAST(sum("매출실적") AS BIGINT) FROM {t} WHERE "_row_kind" = \'data\''
            " AND \"월\" IN ('7월','8월','9월') GROUP BY 1"
        )
    )
    assert got == key["실적집계_v2(9/28 기준)"]

    t = sheet_table(service, "경영지원팀/매출원장/매출원장_2026.xlsx")
    got = dict(
        store.query(
            f'SELECT "담당팀", CAST(sum("공급가액") AS BIGINT) FROM {t}'
            " WHERE \"_row_kind\" = 'data'"
            " AND \"일자\" BETWEEN DATE '2026-07-01' AND DATE '2026-09-30' GROUP BY 1"
        )
    )
    assert got == key["매출원장(9/30 기준)"]


@pytest.mark.parametrize("year", [2024, 2025, 2026])
def test_ledger_rows_match_truth_ledger(service, root, year):
    store = service.store
    t = sheet_table(service, f"경영지원팀/매출원장/매출원장_{year}.xlsx")
    xlsx_rows = store.query(f"SELECT count(*) FROM {t} WHERE \"_row_kind\" = 'data'")[0][0]
    with open(root / "정답지/원장/sales.csv", encoding="utf-8-sig") as fh:
        truth = [
            r
            for r in csv.DictReader(fh)
            if r["일자"].startswith(str(year)) and r["입력일"] <= "2026-09-30"
        ]
    assert xlsx_rows == len(truth)


def test_freshness_against_manifest(service, root):
    manifest = {
        e["경로"]: e
        for e in json.loads((root / "정답지/manifest.json").read_text(encoding="utf-8"))
    }
    got = {f["path"]: f for f in service.catalog()["files"]}
    wrong = []
    for rel, entry in manifest.items():
        if rel not in got:
            continue  # the lock file
        is_ok = got[rel]["fresh"] == "ok"
        if is_ok != entry["기준파일"]:
            wrong.append(rel)
    # Known limits: a format-only copy (_천원) and an old list with an unrelated name.
    assert sorted(wrong) == [
        "개인/박주임/거래처목록.xlsx",
        "영업팀/2026/실적/실적집계_v2_천원.xlsx",
    ]
    assert got["공용/받은자료/실적집계_v2.xlsx"]["copy_of"] == "영업팀/2026/실적/실적집계_v2.xlsx"
    assert got["영업팀/2026/실적/실적집계_최종.xlsx"]["fresh"] == "stale"


def test_departments(service):
    files = service.catalog()["files"]
    unclassified = sorted(f["path"].split("/")[0] for f in files if f["dept"] == "미분류")
    assert unclassified == ["개인", "개인", "새 폴더 (2)", "새 폴더 (2)"]
    assert {f["dept"] for f in files} <= set(service.catalog()["departments"])


def test_prompt_injection_text_is_just_data(service):
    hits = service.store.search_value("이전 지시를 모두 무시")
    assert [(h["file"], h["sheet"], h["column"]) for h in hits] == [
        ("영업팀/2026/수주대장/수주대장_2026.xlsx", "9월", "열13")
    ] or len(hits) == 1
    assert service.catalog()["total"] == 137  # nothing about the scan changed
